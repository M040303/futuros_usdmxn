"""Excel ejecutivo generado desde los resultados exactos de la simulacion."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.chart import LineChart,Reference
from openpyxl.utils import get_column_letter
import pandas as pd
from modelos import calendario
from motor import libro_ruta,estadisticas_por_variable
from formato import formato_excel

NAVY='13253B'; BLUE='1570A6'; LIGHT='EAF2F8'; WHITE='FFFFFF'

# Se comparte con la pestaña de definiciones de Streamlit.
# Cada entrada es (nombre claro, unidad, interpretación).
DEFINICIONES = {
    'modelo': ('Modelo', 'categoría', 'Método con el que se generó el dólar: logística, MBG o ARMA-GARCH-t.'),
    'ruta': ('Ruta', 'identificador', 'Número de escenario dentro de un modelo; no es una fecha ni una observación histórica.'),
    'pnl_bruto': ('Ganancia bruta realizada', 'MXN', 'Suma de las liquidaciones diarias del futuro hasta el cierre efectivo, antes de costos.'),
    'ganancia_neta': ('Ganancia neta realizada', 'MXN', 'Efectivo final menos capital inicial; incluye liquidaciones y costos de entrada y salida.'),
    'efectivo_retirar': ('Efectivo final retirable', 'MXN', 'Caja después de devolver el margen y cobrar costos de salida; puede ser negativa en un escenario extremo.'),
    'capital_inyectado': ('Capital destinado a la operación', 'MXN', 'Margen inicial más costo de entrada y llamadas pagadas. Las llamadas son transferencias desde la caja, no pérdidas adicionales.'),
    'depositos_llamadas': ('Depósitos por llamadas pagadas', 'MXN', 'Suma transferida desde la caja externa hacia el margen antes del cierre efectivo.'),
    'numero_llamadas': ('Número de llamadas pagadas', 'conteo', 'Cantidad de sesiones en las que se pagó una llamada; una llamada impaga que provoca cierre no se cuenta.'),
    'llamada_maxima': ('Mayor llamada pagada', 'MXN', 'Mayor depósito individual al margen realizado en una sesión; cero si no hubo llamadas pagadas.'),
    'dia_llamada_maxima': ('Día de la mayor llamada pagada', 'día natural', 'Días naturales desde la apertura; -1 significa que no hubo llamadas pagadas.'),
    'dolar_llamada_maxima': ('Dólar en la mayor llamada pagada', 'MXN/USD', 'Precio usado como futuro aproximado en el día de la llamada máxima; vacío si no hubo llamada.'),
    'caja_minima': ('Caja externa mínima', 'MXN', 'Menor caja fuera del margen observada durante la operación y al cierre.'),
    'patrimonio_minimo': ('Patrimonio mínimo', 'MXN', 'Menor valor de caja más margen durante la operación y al cierre.'),
    'dia_cierre': ('Día de cierre', 'día natural', 'Días naturales desde la apertura hasta el cierre efectivo, anticipado o programado.'),
    'motivo_cierre': ('Motivo de cierre', 'categoría', 'Programado si llega al día final; sin_liquidez si una llamada no puede pagarse.'),
    'precio_dolar_final': ('Dólar al cierre efectivo', 'MXN/USD', 'Spot simulado del día en que la ruta realmente cierra; puede ser anterior al cierre programado.'),
    'pnl_bruto_horizonte': ('Ganancia bruta hasta el horizonte', 'MXN', 'Ganancia o pérdida hipotética si la posición permaneciera abierta hasta el cierre previsto, sin costos.'),
    'capital_requerido_horizonte': ('Capital requerido hasta el horizonte', 'MXN', 'Margen inicial más costo de entrada y todas las llamadas hipotéticas hasta el cierre previsto; excluye costo de salida.'),
    'depositos_llamadas_horizonte': ('Llamadas hipotéticas acumuladas', 'MXN', 'Suma de llamadas que habría requerido sostener la posición hasta el día final.'),
    'numero_llamadas_horizonte': ('Número de llamadas hipotéticas', 'conteo', 'Sesiones con llamada en la cuenta hipotética que no se cierra por falta de caja.'),
    'llamada_maxima_horizonte': ('Mayor llamada hipotética', 'MXN', 'Mayor depósito individual requerido en la cuenta hipotética hasta el horizonte.'),
    'dia_llamada_maxima_horizonte': ('Día de mayor llamada hipotética', 'día natural', 'Días naturales desde la apertura; -1 significa que no hubo llamada hipotética.'),
    'capital_suficiente_horizonte': ('Capital suficiente hasta el horizonte', '0/1', 'Verdadero cuando el capital requerido hipotético no supera el capital disponible.'),
    'cierre_forzoso': ('Cierre por falta de liquidez', '0/1', 'Verdadero si una llamada impaga obligó a cerrar antes de la fecha programada.'),
}

def crear_excel(c,comparacion,detalle,ejemplos,bandas,estadisticas=None,cierres=None,modo='fijo',aportaciones=None):
    # Compatibilidad con sesiones anteriores de Streamlit durante el despliegue.
    if estadisticas is None or cierres is None:
        resumenes=[]; motivos=[]
        for nombre,grupo in detalle.groupby('modelo',sort=False):
            resumenes.extend(estadisticas_por_variable(grupo,nombre))
            motivos.extend({'modelo':nombre,'motivo_cierre':motivo,'rutas':int(n),
                            'proporcion':float(n/len(grupo))}
                           for motivo,n in grupo.motivo_cierre.value_counts().items())
        estadisticas=pd.DataFrame(resumenes)
        cierres=pd.DataFrame(motivos)
    wb=Workbook(); ws=wb.active; ws.title='Resumen'
    ws.append(['FUTUROS USD/MXN | COMPARACION MONTE CARLO'])
    ws.append(['Escenario academico prospectivo; cotizacion del futuro aproximada por F = spot'])
    ws.append(['Capital (MXN)',c['operacion']['capital'],'Contratos',c['operacion']['contratos_base']])
    ws.append(['Apertura',c['inicio'],'Dias naturales',c['dias']])
    ws.append(['VaR y ES muestrales; cierre forzoso incluido en el P&G neto. Capital requerido: libro sin cierre forzoso.'])
    ws.append([])
    principales=['modelo','simulaciones','capital','prob_capital_suficiente','prob_cierre_forzoso','prob_perdida_neta','mediana_neta','p05_neta','p95_neta','p95_capital_requerido','p99_capital_requerido','VaR95_neto','VaR99_neto','ES95_neto_muestral','ES99_neto_muestral']
    ws.append(principales)
    for row in comparacion[principales].itertuples(index=False,name=None): ws.append([None if str(v)=='nan' else v for v in row])
    st=wb.create_sheet('Estadisticas_variable')
    st.append(['Modelo','Variable','Nombre claro','Unidad','Observaciones','Sin dato',
               'Media','Varianza','Desviación','Mínimo','p05','p25','Mediana',
               'p75','p95','p99','Máximo'])
    for row in estadisticas.itertuples(index=False):
        nombre,unidad,_=DEFINICIONES.get(row.variable,(row.variable,'',''))
        st.append([row.modelo,row.variable,nombre,unidad,row.observaciones,row.sin_dato,
                   row.media,row.varianza,row.desviacion,row.minimo,row.p05,row.p25,
                   row.mediana,row.p75,row.p95,row.p99,row.maximo])
    fr=wb.create_sheet('Motivos_cierre')
    fr.append(['Modelo','Motivo','Rutas','Proporción'])
    for row in cierres.itertuples(index=False): fr.append(list(row))
    dic=wb.create_sheet('Diccionario')
    dic.append(['Variable','Nombre claro','Unidad','Definición'])
    for variable,(nombre,unidad,descripcion) in DEFINICIONES.items():
        dic.append([variable,nombre,unidad,descripcion])
    md=wb.create_sheet('Parametros')
    md.append(['PARAMETRO','VALOR'])
    def flat(x,prefix=''):
        for k,v in x.items():
            key=f'{prefix}.{k}' if prefix else k
            if isinstance(v,dict): yield from flat(v,key)
            else: yield key,str(v) if isinstance(v,(list,tuple)) else v
    for row in flat(c): md.append(list(row))
    notes=wb.create_sheet('Notas')
    for line in [
        'Contrato DA: USD 10,000 por contrato. Verificar ficha y tarifa vigente antes de operar.',
        'Futuro proxy: F_t = redondeo(S_t/tick)*tick; no incorpora base, carry ni diferencial de tasas.',
        'MXN 50 por lado y MXN 15,000/11,000 de margen son supuestos del ejercicio, no tarifas o garantias oficiales vigentes.',
        'Modelo logistico: niveles iid; se ajusta a precios historicos, no a rendimientos.',
        'MBG: S_t=S_(t-1)*exp(media_log + sigma*Z_t).',
        'GARCH: ARMA(2,2) para media; h_t=omega+alpha*e_(t-1)^2+beta*h_(t-1); innovacion t estandarizada.',
        'En GARCH-t, la esperanza teorica del precio exponenciado puede ser infinita: comparar medianas y cuantiles.',
        'Margen = margen previo + P&G diario. Si baja del mantenimiento, reponer hasta objetivo.',
        'Si no hay caja para la llamada: cierre forzoso al precio de la sesion y pago de costos de salida.',
        'Capital requerido se calcula siguiendo el libro hipotetico hasta cierre sin limitar la caja.',
        'VaR 95/99 = cuantiles 95/99 de la perdida neta (-P&G neto). ES = promedio muestral sobre el cuantil.',
        'Estas probabilidades dependen de supuestos, ventana de calibracion, datos y semilla. No son garantias.',
        'No se incluyen intereses, impuestos, cambios de margen, liquidez del futuro ni liquidacion fisica.'
    ]: notes.append([line])
    m=wb.create_sheet('Rutas_ejemplo'); m.append(list(detalle.columns))
    for modelo in detalle.modelo.unique():
        for row in detalle[detalle.modelo==modelo].head(1).itertuples(index=False,name=None):
            m.append([None if str(v)=='nan' else v for v in row])
    notes.append(['Estadisticas_variable usa todas las rutas de cada modelo; Rutas_ejemplo muestra una ruta por modelo para auditoría.'])
    fechas,_=calendario(c)
    tr=wb.create_sheet('Trayectorias_muestra')
    tr.append(['Modelo','Ruta_muestra','Fecha','Spot_MXN_USD'])
    for modelo,paths in ejemplos.items():
        for i,path in enumerate(paths,1):
            for d,spot in zip(fechas,path): tr.append([modelo,i,d.isoformat(),float(spot)])
    for modelo,paths in ejemplos.items():
        if len(paths):
            ledger=libro_ruta(paths[0],c,modo=modo)
            sheet=wb.create_sheet('Flujo_'+modelo[:20]);sheet.append(list(ledger.columns))
            for row in ledger.itertuples(index=False,name=None): sheet.append(list(row))
    b=wb.create_sheet('Bandas_muestra');b.append(['Modelo','Dia','p05','Mediana','p95'])
    for row in bandas.itertuples(index=False): b.append(list(row))
    if aportaciones is not None:
        ap=wb.create_sheet('Aportaciones_fecha')
        ap.append(list(aportaciones.columns))
        for row in aportaciones.itertuples(index=False,name=None):
            ap.append(list(row))
        notes.append(['Aportaciones_fecha: percentiles exactos por fecha sobre todas las simulaciones. Cuenta hipotética sin cierre forzoso. No incluye margen inicial ni costos.'])
        notes.append(['P95 por fecha no es una trayectoria ni una banda simultánea del 95 %. No sumar percentiles diarios. El P95 acumulado final dimensiona la reserva de llamadas de todo el periodo.'])
    for sheet in wb:
        sheet.freeze_panes='B2'; sheet.sheet_view.showGridLines=False
        for cell in sheet[1]: cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(bold=True,color=WHITE,size=11)
        sheet.row_dimensions[1].height=24
        for col in sheet.columns:
            letter=get_column_letter(col[0].column)
            maxlen=max((len(str(v.value)) for v in list(col)[:1000] if v.value is not None),default=10)
            sheet.column_dimensions[letter].width=min(max(13,maxlen+2),70)
    for cell in ws[7]: cell.fill=PatternFill('solid',fgColor=BLUE);cell.font=Font(bold=True,color=WHITE)
    # Formato nativo: los importes siguen siendo números editables en Excel.
    for sheet in wb:
        header_row = 7 if sheet.title == 'Resumen' else 1
        headers = [cell.value for cell in sheet[header_row]]
        for row in sheet.iter_rows(min_row=header_row+1):
            for cell, header in zip(row, headers):
                fmt = formato_excel(header)
                if sheet.title == 'Estadisticas_variable':
                    fmt = None
                    if cell.column >= 7:
                        variable = sheet.cell(cell.row, 2).value
                        fmt = formato_excel(variable)
                        if cell.column == 8:
                            unidad = sheet.cell(cell.row, 4).value
                            fmt = '#,##0.00 "MXN²"' if unidad == 'MXN' else '#,##0.0000 "(MXN/USD)²"' if unidad == 'MXN/USD' else '0.0000'
                        elif variable in ('capital_suficiente_horizonte', 'cierre_forzoso'):
                            fmt = '0.00%'
                if sheet.title == 'Parametros' and cell.column == 2:
                    key = str(sheet.cell(cell.row, 1).value)
                    if key.endswith(('margen_fraccion','mantenimiento_fraccion','objetivo_fraccion','media_log_diaria','sigma_diaria','mu_diaria')):
                        fmt = '0.0000%'
                    elif key.endswith(('capital','margen_inicial','margen_mantenimiento','objetivo_reposicion','comision_entrada','comision_salida')):
                        fmt = formato_excel('capital')
                    elif key in ('s0','logistica.location','logistica.scale','operacion.tick','operacion.deslizamiento_por_lado'):
                        fmt = formato_excel('spot')
                if fmt and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = fmt
    ws['B3'].number_format = '"$"#,##0.00'
    wb.active=0
    stream=BytesIO();wb.save(stream);return stream.getvalue()


