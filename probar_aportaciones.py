"""Conciliación de llamadas por fecha, cuantiles y compatibilidad."""
from copy import deepcopy
from io import BytesIO
import numpy as np
from openpyxl import load_workbook
from modelos import cargar, generar, calendario
from motor import procesar_lote, libro_ruta, simular, parametros
from excel import crear_excel
from aportaciones import graficas_aportaciones

def probar():
    c=cargar()
    fechas,sesiones=calendario(c)
    for modo in ('fijo','porcentaje'):
        for direccion in (-1,1):
            cfg=deepcopy(c);cfg['operacion']['direccion']=direccion
            paths=generar(cfg,'gbm',30,39)
            resultado,diario=procesar_lote(paths,cfg,3,modo,con_aportaciones=True)
            # Oráculo independiente: libro sin falta de caja, una ruta a la vez.
            ilimitado=deepcopy(cfg);ilimitado['operacion']['capital']=1e12
            for i,path in enumerate(paths):
                libro=libro_ruta(path,ilimitado,3,modo)
                np.testing.assert_allclose(diario[i],libro.llamada,atol=1e-8)
            np.testing.assert_allclose(diario.sum(axis=1),resultado.depositos_llamadas_horizonte)
            assert not diario[:,~sesiones].any()
            assert not diario[:,[0,-1]].any()
    # Ruta con cierre forzoso: no debe censurarse su necesidad posterior.
    cfg=deepcopy(c);cfg['operacion']['capital']=47000.
    path=np.full((1,c['dias']+1),c['s0'])
    for j in range(1,len(fechas)): path[0,j]=c['s0']-.2*j
    resultado,diario=procesar_lote(path,cfg,3,'fijo',con_aportaciones=True)
    assert resultado.cierre_forzoso.iloc[0]
    assert diario[0,int(resultado.dia_cierre.iloc[0])+1:].sum()>0
    # Más de 20,000 rutas: confirma que no se usa el detalle recortado.
    salida=simular(c,['gbm'],20507,5000,42,con_estadisticas=True,con_aportaciones=True)
    comp,detalle,ejemplos,bandas,stats,cierres,aportes=salida
    assert len(detalle)==20000 and (aportes.simulaciones==20507).all()
    assert np.all(np.diff(aportes.aporte_acumulado_p95)>=-1e-8)
    assert np.all(np.diff(aportes.prob_reserva_suficiente)<=1e-12)
    ultimo=aportes.iloc[-1];mi,mt,obj,ce,cs=parametros(c,3,'fijo')
    assert np.isclose(ultimo.aporte_acumulado_p95+mi+ce,comp.p95_capital_requerido.iloc[0])
    assert np.isclose(ultimo.prob_reserva_suficiente,comp.prob_capital_suficiente.iloc[0])
    acumulados=[];diarios=[]
    for num,done in enumerate(range(0,20507,5000)):
        paths=generar(c,'gbm',min(5000,20507-done),42+num)
        _,d=procesar_lote(paths,c,3,'fijo',con_aportaciones=True)
        diarios.append(d);acumulados.append(np.cumsum(d,axis=1))
    np.testing.assert_allclose(aportes.aporte_acumulado_p95,np.quantile(np.vstack(acumulados),.95,axis=0))
    np.testing.assert_allclose(aportes.aporte_diario_p95,np.quantile(np.vstack(diarios),.95,axis=0))
    plano=deepcopy(c);plano['gbm']['sigma_diaria']=0.;plano['gbm']['media_log_diaria']=0.
    cero=simular(plano,['gbm'],20,20,con_aportaciones=True)[-1]
    assert not cero.aporte_acumulado_p95.any() and (cero.prob_reserva_suficiente==1).all()
    wb=load_workbook(BytesIO(crear_excel(c,comp,detalle,ejemplos,bandas,stats,cierres,aportaciones=aportes)))
    sheet=wb['Aportaciones_fecha'];idx=list(aportes.columns).index('aporte_acumulado_p95')+1
    assert sheet.max_row==len(fechas)+1 and '$' in sheet.cell(2,idx).number_format
    lineas,barras=graficas_aportaciones(aportes)
    assert len(lineas.data)==2 and len(barras.data)==1
    assert len(lineas.layout.shapes)==1
    assert len(simular(c,['gbm'],20,20))==4
    assert len(simular(c,['gbm'],20,20,con_estadisticas=True))==6
    print('Aprobadas: libro diario independiente, largos/cortos, ambos márgenes, cierre forzoso, días sin sesión, 20,507 rutas completas, cuantiles exactos, gráficas, Excel y compatibilidad.')

if __name__=='__main__': probar()
